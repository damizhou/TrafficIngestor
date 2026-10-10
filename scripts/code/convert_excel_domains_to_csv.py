#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""将 Excel 的 domain 列转换为 URL 收集器使用的 id,url,domain CSV。"""

from __future__ import annotations

import argparse
import csv
import re
from pathlib import Path

from openpyxl import load_workbook


SCRIPTS_DIR = Path(__file__).resolve().parent.parent
DEFAULT_INPUT = SCRIPTS_DIR / "origin" / "subpages.xlsx"
DEFAULT_OUTPUT = SCRIPTS_DIR / "subpages_seeds.csv"
DOMAIN_LABEL = re.compile(r"[a-z0-9](?:[a-z0-9-]{0,61}[a-z0-9])?")


def read_domains(input_path: Path, sheet_name: str) -> list[str]:
    workbook = load_workbook(input_path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"工作表不存在：{sheet_name}，可用工作表：{workbook.sheetnames}")
        records = workbook[sheet_name].iter_rows(values_only=True)
        header = next(records, ())
        columns = [str(value).strip().lower() for value in header]
        if columns.count("domain") != 1:
            raise ValueError("首行必须包含且仅包含一个 domain 列")
        column = columns.index("domain")
        domains = []
        seen = set()
        for row_number, row in enumerate(records, start=2):
            value = row[column]
            if value is None or not str(value).strip():
                continue
            domain = str(value).strip().lower()
            labels = domain.split(".")
            if (len(domain) > 253 or len(labels) < 2
                    or any(DOMAIN_LABEL.fullmatch(label) is None for label in labels)):
                raise ValueError(f"第 {row_number} 行不是有效的裸域名：{value!r}")
            if domain not in seen:
                seen.add(domain)
                domains.append(domain)
        if not domains:
            raise ValueError("Excel 中没有可用域名")
        return domains
    finally:
        workbook.close()


def convert(input_path: Path, output_path: Path, sheet_name: str = "Sheet1") -> int:
    input_path = input_path.expanduser().resolve()
    output_path = output_path.expanduser().resolve()
    if input_path == output_path:
        raise ValueError("输入与输出文件不能相同")
    if output_path.suffix.lower() != ".csv":
        raise ValueError("输出文件必须使用 .csv 扩展名")
    domains = read_domains(input_path, sheet_name)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary = output_path.with_suffix(".csv.tmp")
    try:
        with temporary.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(["id", "url", "domain"])
            writer.writerows(
                (index, f"https://{domain}", domain)
                for index, domain in enumerate(domains, start=1)
            )
        temporary.replace(output_path)
    finally:
        temporary.unlink(missing_ok=True)
    return len(domains)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT, help="输入 Excel 路径")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT, help="输出 CSV 路径，覆盖同名文件")
    parser.add_argument("--sheet", default="Sheet1", help="工作表名称，默认 Sheet1")
    args = parser.parse_args()
    count = convert(args.input, args.output, args.sheet)
    print(f"已生成：{args.output.expanduser().resolve()}，共 {count} 个唯一域名")


if __name__ == "__main__":
    main()
